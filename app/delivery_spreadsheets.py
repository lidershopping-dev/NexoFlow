"""Importação privada dos exports de conversas e contatos do Delivery."""

from __future__ import annotations

import ast
import hashlib
import logging
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from uuid import UUID
from zoneinfo import ZoneInfo

from delivery_tickets import delivery_ticket_key, is_order_tag
from farmalider_integration import _identity_name
from urn_utils import native_contact_urns

CONVERSATIONS_FILENAME = "Conversas_Delivery.xlsx"
CONTACTS_FILENAME = "Contatos_Delivery.xlsx"
EXPORT_TIMEZONE = ZoneInfo("America/Belem")


class DeliverySpreadsheetError(ValueError):
    """Erro seguro: não inclui nomes, CPFs nem conteúdo das células."""


def _rows(path, sheet, required):
    import openpyxl

    path = Path(path)
    if not path.is_file():
        raise DeliverySpreadsheetError(
            f"Planilha ausente: use {CONVERSATIONS_FILENAME} e {CONTACTS_FILENAME} na pasta importacoes."
        )
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 100 * 1024 * 1024:
        raise DeliverySpreadsheetError("Use uma planilha .xlsx de até 100 MB.")
    try:
        workbook = openpyxl.load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
    except Exception:
        raise DeliverySpreadsheetError(
            "Não foi possível abrir o export XLSX; verifique a integridade do arquivo."
        ) from None
    try:
        sheets = (
            [
                name
                for name in workbook.sheetnames
                if re.fullmatch(r"Contact(?: [1-9][0-9]*)?", name)
            ]
            if sheet == "Contact"
            else [sheet]
            if sheet in workbook.sheetnames
            else []
        )
        if not sheets:
            raise DeliverySpreadsheetError(
                f"Aba obrigatória ausente: {sheet}. Preserve a estrutura original do export."
            )
        for sheet_name in sheets:
            worksheet = workbook[sheet_name]
            # Weni's contact export omits worksheet dimensions in its XML.
            worksheet.reset_dimensions()
            iterator = worksheet.iter_rows()
            header = next(iterator, ())
            if any(cell.data_type == "f" for cell in header):
                raise DeliverySpreadsheetError(
                    "O export contém fórmulas no cabeçalho; preserve os cabeçalhos originais."
                )
            names = [str(cell.value or "").strip() for cell in header]
            if not required <= set(names):
                raise DeliverySpreadsheetError(
                    "Colunas obrigatórias ausentes; preserve os cabeçalhos originais dos exports."
                )
            for number, cells in enumerate(iterator, 2):
                values = []
                for index, cell in enumerate(cells):
                    value = cell.value
                    if cell.data_type == "f":
                        # The original Chats export serializes boolean constants
                        # as formulas. Decode only these exact constants in the
                        # activity column, without evaluating Excel expressions
                        # or trusting potentially stale cached formula values.
                        if (
                            sheet == "rooms"
                            and index < len(names)
                            and names[index] == "is_active"
                            and isinstance(value, str)
                            and value.upper() in ("=FALSE()", "=TRUE()")
                        ):
                            value = value.upper() == "=TRUE()"
                        else:
                            raise DeliverySpreadsheetError(
                                "O export contém uma fórmula não permitida "
                                f"na célula {cell.coordinate}; use o arquivo original de dados."
                            )
                    values.append(value)
                yield number, names, values
    finally:
        workbook.close()


def _uuid(value):
    try:
        return str(UUID(str(value or "").strip()))
    except ValueError:
        return ""


def _date(value):
    try:
        date = (
            value
            if isinstance(value, datetime)
            else datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        )
    except (ValueError, TypeError):
        return None
    # The Chats exporter converts timestamps to the project's timezone and
    # strips tzinfo. This import explicitly targets the Delivery Belem project.
    return date.replace(tzinfo=EXPORT_TIMEZONE) if date.tzinfo is None else date


def _tags(value):
    if not isinstance(value, str) or len(value) > 10000:
        return []
    try:
        parsed = ast.literal_eval(value)
    except (SyntaxError, ValueError, RecursionError):
        return []
    return (
        parsed
        if isinstance(parsed, list) and all(isinstance(tag, str) for tag in parsed)
        else []
    )


def read_delivery_spreadsheets(directory, *, cutoff):
    """Map Chats UUIDs to Flows UUIDs via the existing UUID/native URN link."""
    directory = Path(directory)
    contacts, urn_index, summary, rejected = {}, defaultdict(set), Counter(), []
    contact_path, room_path = (
        directory / CONTACTS_FILENAME,
        directory / CONVERSATIONS_FILENAME,
    )
    aliases = {"Pedido Produtos": "pedido_produtos", "ID do Cliente": "id_do_cliente"}
    for _number, headers, values in _rows(
        contact_path, "Contact", {"Contact UUID", "Name", "Field:CPF", "Field:Email"}
    ):
        row = dict(zip(headers, values, strict=False))
        uid = _uuid(row.get("Contact UUID"))
        if not uid:
            summary["CONTATO_UUID_INVALIDO"] += 1
            continue
        fields, urns = {}, []
        for header, value in zip(headers, values, strict=False):
            if header.startswith("Field:"):
                key = header[6:]
                fields[aliases.get(key, key.lower().replace(" ", "_"))] = (
                    "" if value is None else str(value)
                )
            elif header.startswith("URN:") and value:
                scheme = header[4:].lower()
                if scheme == "deleted":
                    continue
                address = str(value).strip()
                urns.append(
                    address
                    if address.startswith(scheme + ":")
                    else scheme + ":" + address
                )
        contact = {
            "uuid": uid,
            "name": str(row.get("Name") or ""),
            "fields": fields,
            "urns": urns,
        }
        if uid in contacts and contacts[uid] != contact:
            raise DeliverySpreadsheetError(
                "UUID de contato duplicado com dados conflitantes; gere um novo export de contatos."
            )
        contacts[uid] = contact
        for urn in urns:
            urn_index[urn].add(uid)
    tickets, seen = [], set()
    required = {
        "contact__uuid",
        "contact__name",
        "urn",
        "created_on",
        "ended_at",
        "is_active",
        "tags_list",
        "user__email",
    }
    for number, headers, values in _rows(room_path, "rooms", required):
        row = dict(zip(headers, values, strict=False))
        tags = _tags(row.get("tags_list"))
        if not any(is_order_tag(tag) for tag in tags):
            continue
        summary["ATENDIMENTOS_COM_TAG"] += 1
        reason = ""
        sector = row.get("queue__sector__name")
        if sector and _identity_name(sector) != "delivery farmalider":
            reason = "SETOR_DIVERGENTE"
        opened, ended = _date(row.get("created_on")), _date(row.get("ended_at"))
        if (
            not opened
            or not ended
            or opened > ended
            or ended > cutoff
            or row.get("is_active") not in (False, 0, "False", "false")
        ):
            reason = reason or "ENCERRAMENTO_AUSENTE_OU_FORA_DO_CORTE"
        if not row.get("user__email"):
            reason = reason or "ATENDENTE_NAO_IDENTIFICADO"
        chat_uid, urn = (
            _uuid(row.get("contact__uuid")),
            str(row.get("urn") or "").strip(),
        )
        ids = urn_index.get(urn, set())
        if chat_uid in contacts:
            ids = {chat_uid} if not urn or chat_uid in ids else set()
        if not chat_uid or len(ids) != 1:
            reason = reason or (
                "URN_AMBIGUA" if len(ids) > 1 else "VINCULO_WENI_NAO_LOCALIZADO"
            )
        uid = next(iter(ids)) if len(ids) == 1 else ""
        if reason or opened is None or ended is None:
            summary[reason] += 1
            rejected.append(
                {"linha": number, "chat_contact_uuid": chat_uid, "motivo": reason}
            )
            continue
        key = delivery_ticket_key(chat_uid, opened.isoformat(), ended.isoformat())
        if key in seen:
            summary["ATENDIMENTO_DUPLICADO"] += 1
            continue
        seen.add(key)
        tickets.append(
            {
                "ticket_uuid": key,
                "chats_room_uuid": "",
                "chat_contact_uuid": chat_uid,
                "origem_registro": "PLANILHAS",
                "protocolo": str(row.get("protocol") or ""),
                "contact_uuid": uid,
                "tags": ", ".join(tags),
                "aberto_em": opened.isoformat(),
                "encerrado_em": ended.isoformat(),
                "custom_fields": {},
                "export_urn": urn,
                "export_contact_name": str(row.get("contact__name") or ""),
            }
        )
    selected_ids = {ticket["contact_uuid"] for ticket in tickets}
    summary.update(
        contatos_export=len(contacts),
        tickets_selecionados=len(tickets),
        clientes_vinculados=len(selected_ids),
    )
    sources = [
        {"arquivo": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        for path in (room_path, contact_path)
    ]
    return (
        tickets,
        [contacts[uid] for uid in sorted(selected_ids)],
        {"summary": dict(summary), "rejections": rejected, "sources": sources},
    )


def validate_current_contacts(tickets, current_contacts, *, refresh_contact=None):
    """Require the exported UUID/native URN link to identify one live contact."""
    by_uuid, index = {}, defaultdict(set)
    duplicates, conflicts = 0, set()
    for contact in current_contacts:
        uid = str(contact.get("uuid") or "")
        if not uid:
            raise DeliverySpreadsheetError(
                "Contato sem UUID na API; atualização interrompida."
            )
        if uid in by_uuid:
            duplicates += 1
            if contact != by_uuid[uid]:
                conflicts.add(uid)
        else:
            by_uuid[uid] = contact
        for urn in native_contact_urns(contact):
            index[urn].add(uid)
    for uid in sorted(conflicts):
        if refresh_contact is None:
            raise DeliverySpreadsheetError(
                "UUID repetido com dados diferentes; releitura individual necessária antes da atualização."
            )
        fresh = refresh_contact(uid)
        if fresh is None:
            by_uuid.pop(uid, None)
            continue
        if str(fresh.get("uuid") or "") != uid:
            raise DeliverySpreadsheetError(
                "Releitura devolveu outro UUID; atualização interrompida."
            )
        by_uuid[uid] = fresh
        for urn in native_contact_urns(fresh):
            index[urn].add(uid)
    if duplicates:
        logging.warning(
            "Weni paginação: repetições=%s; UUIDs únicos=%s; releituras por divergência=%s. Ambiguidades de identidade permanecem bloqueadas.",
            duplicates,
            len(by_uuid),
            len(conflicts),
        )
    accepted, rejected = [], Counter()
    for ticket in tickets:
        uid, urn = ticket["contact_uuid"], ticket.get("export_urn")
        current = by_uuid.get(uid)
        if current is None:
            rejected["UUID_NAO_LOCALIZADO_NA_WENI_ATUAL"] += 1
        elif urn and (
            urn not in native_contact_urns(current) or index.get(urn, set()) != {uid}
        ):
            rejected["URN_ALTERADA_OU_AMBIGUA_NA_WENI_ATUAL"] += 1
        else:
            accepted.append(ticket)
    selected_ids = {ticket["contact_uuid"] for ticket in accepted}
    return accepted, [by_uuid[uid] for uid in sorted(selected_ids)], dict(rejected)
